import base64
from types import SimpleNamespace

import pytest
from pydantic import ValidationError

from agentai.chat.attachments import AttachmentStore, SOURCE, extract_attachments
from agentai.chat.service import run_chat
from agentai.commands import Attachment, Chat
from agentai.errors import ApplicationError


USER = {'name': 'reader', 'roles': ['reader'], 'sources': []}


def file(name='notes.txt', text='Проверочный номер договора: ZM-7284.'):
    return Attachment(filename=name, content_base64=base64.b64encode(text.encode()).decode())


def test_reader_attachment_uses_private_store_without_touching_shared_search():
    shared = SimpleNamespace(search=lambda *args: pytest.fail('Shared search must not run'))
    def answer(question, user, store, model, **options):
        rows = store.search(question, user['sources'])
        assert 'ZM-7284' in rows[0]['text']
        assert store.read_source(rows[0]['id'], user['sources'], 0, 6) == rows
        return {'answer': 'ZM-7284', 'sources': rows}
    result = run_chat(Chat(question='Какой номер договора?', attachments=[file()]), USER,
                      services=lambda: (shared, None), answer=answer)
    assert result['attachments'][0]['filename'] == 'notes.txt'
    assert result['attachment_rows'][0]['attachment']
    assert USER['sources'] == []


def test_same_named_files_are_distinct_and_neighbors_do_not_cross_files():
    rows, _ = extract_attachments([file(text='a\n' * 100), file(text='b\n' * 100)])
    assert len({r['document_id'] for r in rows}) == 2
    store = AttachmentStore(rows)
    neighbors = store.read_source(rows[0]['id'], [SOURCE], 0, 6)
    assert all(r['document_id'] == rows[0]['document_id'] for r in neighbors)
    assert not store.search('номер', [])
    assert not store.read_source(rows[0]['id'], [], 0, 6)


@pytest.mark.parametrize('item', [Attachment(filename='file.txt', content_base64='!bad!'),
    file('bad.exe'), file('empty.txt', ''), file('invalid\n.txt')])
def test_invalid_files_have_clear_errors(item):
    with pytest.raises(ApplicationError) as error:
        extract_attachments([item])
    assert error.value.status_code == 422


def test_limits_before_queueing():
    with pytest.raises(ValidationError):
        Chat(question='test', attachments=[file()] * 4)
    with pytest.raises(ValidationError):
        Chat(question='test', attachments=[Attachment(filename='large.txt', content_base64='a' * 20_000_001)])


def test_code_files_are_extracted_without_renaming():
    rows, _ = extract_attachments([file('logic.py', 'def retries():\n    return 7')])
    assert rows[0]['citation'].endswith('/logic.py#L1-L2')
    assert 'return 7' in rows[0]['text']


def test_docx_paragraphs_and_tables_are_available():
    from io import BytesIO
    from docx import Document
    doc = Document()
    doc.add_paragraph('Контракт ZM-7284')
    table = doc.add_table(rows=2, cols=1)
    table.cell(0, 0).text = 'Попытки'
    table.cell(1, 0).text = '7'
    stream = BytesIO()
    doc.save(stream)
    rows, _ = extract_attachments([Attachment(filename='contract.docx', content_base64=base64.b64encode(stream.getvalue()).decode())])
    assert any('ZM-7284' in row['text'] for row in rows)
    assert any('Попытки: 7' in row['text'] for row in rows)


def test_pdf_text_extraction_and_malformed_pdf(monkeypatch):
    from agentai.ingestion import pdf
    monkeypatch.setattr(pdf, 'pages', lambda data: iter(['Contract ZM-7284', 'Attempts: 7']))
    rows, _ = extract_attachments([Attachment(filename='contract.pdf', content_base64=base64.b64encode(b'fixture').decode())])
    assert '?page=2' in rows[-1]['citation'] and rows[-1]['text'] == 'Attempts: 7'
    def broken(data):
        raise ValueError('parser internals must not leak')
    monkeypatch.setattr(pdf, 'pages', broken)
    with pytest.raises(ApplicationError) as error:
        extract_attachments([file('bad.pdf')])
    assert 'bad.pdf' in error.value.detail and 'internals' not in error.value.detail


def test_explicit_followup_reuses_file_but_new_topic_does_not():
    rows, _ = extract_attachments([file()])
    class Memory:
        def get(self, *args):
            return {'turns': [{'question': 'Первый вопрос', 'answer': 'Ответ', 'sources': rows,
                              'context': {'attachment_rows': rows}}]}
        def append(self, *args, **kwargs):
            pass
    shared = object()
    def answer(question, user, store, model, **kwargs):
        expected = AttachmentStore if 'файле' in question else type(shared)
        assert isinstance(store, expected)
        return {'answer': 'ok', 'sources': []}
    import uuid
    for question in ['Что ещё в этом файле?', 'Как создать отчёт?']:
        run_chat(Chat(question=question, conversation_id=uuid.uuid4()), USER,
                 services=lambda: (shared, None), answer=answer, memory=Memory())
