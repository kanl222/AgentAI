import base64
import hashlib
import json
from io import BytesIO
from unittest.mock import Mock
import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from qdrant_client import QdrantClient
from langchain_core.messages import AIMessage
from agentai import api
from agentai.core import Store, chunks, document, authenticate, require
from agentai.agent import answer, skills


class Embeddings:
    def embed_documents(self,texts):
        return [[1.,0.,0.] for text in texts]
    def embed_query(self,text):
        return [1.,0.,0.]


@pytest.fixture
def store():
    return Store(Embeddings(),QdrantClient(':memory:'))


@pytest.fixture
def client(monkeypatch,store):
    users = [{'name':'alice','sha256':hashlib.sha256(b'alice-key').hexdigest(),'roles':['reader'],'sources':['a']},
             {'name':'admin','sha256':hashlib.sha256(b'admin-key').hexdigest(),'roles':['reader','indexer'],'sources':['a']}]
    monkeypatch.setenv('AGENT_USERS',json.dumps(users))
    monkeypatch.setattr(api,'services',lambda:(store,Mock()))
    return TestClient(api.app)


def test_authentication_and_authorization(client):
    assert client.post('/chat',json={'question':'hello'}).status_code==401
    assert client.get('/v1/models',headers={'Authorization':'Bearer bad'}).status_code==401
    body={'source':'a','filename':'test.md','content_base64':base64.b64encode(b'hello').decode()}
    assert client.post('/ingest/document',json=body,headers={'Authorization':'Bearer alice-key'}).status_code==403
    body['source']='private'
    assert client.post('/ingest/document',json=body,headers={'Authorization':'Bearer admin-key'}).status_code==403


def test_upload_and_acl(client,store):
    body={'source':'a','filename':'test.md','content_base64':base64.b64encode(b'platform owner').decode()}
    assert client.post('/ingest/document',json=body,headers={'Authorization':'Bearer admin-key'}).json()['chunks']==1
    store.replace('secret',chunks('private password','secret','hidden'))
    assert all(r['source']=='a' for r in store.search('password',['a']))
    assert store.search('owner',[])==[]
    assert 'private password' not in json.dumps(store.search('owner',['a']))


def test_reindex_and_lines(store):
    store.replace('a',chunks('old text','a','gitlab:1@'+'a'*40+'/x.py'))
    store.replace('a',chunks('new text','a','gitlab:1@'+'b'*40+'/x.py'))
    rows=store.search('new',['a'])
    assert len(rows)==1 and rows[0]['text']=='new text'
    assert rows[0]['citation'].endswith('/x.py#L1-L1')
    assert 'b'*40 in rows[0]['citation']


def test_document_locations():
    from docx import Document
    doc=Document(); doc.add_paragraph('hello')
    table=doc.add_table(rows=1,cols=1);table.cell(0,0).text='world'
    out=BytesIO();doc.save(out)
    rows=list(document(out.getvalue(),'demo.docx','a'))
    assert '?paragraph=1' in rows[0]['citation']
    assert '?table=1&row=1' in rows[1]['citation']
    assert list(document(b'a\nb','a.txt','a'))[0]['citation'].endswith('#L1-L2')
    with pytest.raises(ValueError):list(document(b'bad','a.exe','a'))


def test_skills_catalog():
    assert {'documents','code'}<=skills().keys()


class LoopModel:
    def bind_tools(self,tools):return self
    def invoke(self,messages):
        return AIMessage(content='',tool_calls=[{'name':'search','args':{'query':'owner'},'id':'call-'+str(len(messages))}])


def test_graph_limit_and_evidence(store):
    store.replace('a',chunks('owner Platform','a','demo.md'))
    result=answer('owner?',{'roles':['reader'],'sources':['a']},store,LoopModel(),max_steps=2)
    assert result['steps']==2
    assert 'лимит' in result['answer']
    assert result['sources'][0]['citation']=='demo.md#L1-L1'


def test_openai_stream_and_health(client):
    assert client.get('/health').status_code==200
    assert client.post('/v1/chat/completions',json={'messages':[],'stream':True},headers={'Authorization':'Bearer alice-key'}).status_code==400


def test_sql_role_denied():
    with pytest.raises(HTTPException):require({'roles':['reader'],'sources':[]},'sql_reader')


def test_hybrid_lexical_boost(store):
    for i in range(6):
        store.replace(str(i),chunks('needle identifier' if i==5 else 'ordinary filler',str(i),str(i)))
    result=store.search('needle',[str(i) for i in range(6)])
    assert result[0]['source']=='5'
